"""
ED1 — PowerPoint visual: 3 matrizes alinhadas (R,G,B) no inicio.
Gera Matrizes_e_Sistemas_Lineares_Imagens_VISUAL.pptx

Uso:
  python app-python.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw

OUT_DIR = Path(__file__).resolve().parent
ASSETS = OUT_DIR / "_ed1_assets"
ASSETS.mkdir(exist_ok=True)

R = [[196, 184, 201], [110, 85, 89], [38, 33, 29]]
G = [[188, 171, 189], [85, 56, 58], [41, 39, 39]]
B = [[177, 162, 177], [80, 52, 55], [32, 29, 28]]

COLOR_BG = RGBColor(248, 249, 250)
COLOR_NAVY = RGBColor(15, 32, 67)
COLOR_BLUE = RGBColor(30, 114, 219)
COLOR_DARK = RGBColor(33, 37, 41)
COLOR_MUTED = RGBColor(108, 117, 125)
COLOR_CARD = RGBColor(255, 255, 255)
COLOR_LINE = RGBColor(220, 224, 230)
COLOR_SOFT = RGBColor(230, 242, 255)

CELL = 72
GAP = 4


def transpose(m):
    return [list(col) for col in zip(*m)]


def scale_mat(m, k):
    return [[max(0, min(255, int(round(v * k)))) for v in row] for row in m]


def draw_grid_rgb(path: Path, r, g, b):
    n = len(r)
    w = n * CELL + (n + 1) * GAP
    img = Image.new("RGB", (w, w), (248, 249, 250))
    draw = ImageDraw.Draw(img)
    for i in range(n):
        for j in range(n):
            x0 = GAP + j * (CELL + GAP)
            y0 = GAP + i * (CELL + GAP)
            color = (r[i][j], g[i][j], b[i][j])
            draw.rounded_rectangle(
                [x0, y0, x0 + CELL - 1, y0 + CELL - 1],
                radius=10,
                fill=color,
                outline=(30, 114, 219),
                width=2,
            )
            lum = 0.299 * color[0] + 0.587 * color[1] + 0.114 * color[2]
            tc = (255, 255, 255) if lum < 140 else (20, 20, 20)
            txt = str(color[0])
            bbox = draw.textbbox((0, 0), txt)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            draw.text((x0 + (CELL - tw) / 2, y0 + (CELL - th) / 2), txt, fill=tc)
    img.save(path)
    return path


def draw_grid_gray(path: Path, m, tint=None):
    n = len(m)
    w = n * CELL + (n + 1) * GAP
    img = Image.new("RGB", (w, w), (248, 249, 250))
    draw = ImageDraw.Draw(img)
    for i in range(n):
        for j in range(n):
            v = m[i][j]
            if tint == "R":
                color = (v, 40, 40)
            elif tint == "G":
                color = (40, v, 40)
            elif tint == "B":
                color = (40, 40, v)
            else:
                color = (v, v, v)
            x0 = GAP + j * (CELL + GAP)
            y0 = GAP + i * (CELL + GAP)
            draw.rounded_rectangle(
                [x0, y0, x0 + CELL - 1, y0 + CELL - 1],
                radius=10,
                fill=color,
                outline=(30, 114, 219),
                width=2,
            )
            lum = 0.299 * color[0] + 0.587 * color[1] + 0.114 * color[2]
            tc = (255, 255, 255) if lum < 140 else (20, 20, 20)
            txt = str(v)
            bbox = draw.textbbox((0, 0), txt)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            draw.text((x0 + (CELL - tw) / 2, y0 + (CELL - th) / 2), txt, fill=tc)
    img.save(path)
    return path


def draw_swatch(path: Path, rgb, caption: str):
    img = Image.new("RGB", (220, 220), (248, 249, 250))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([10, 10, 210, 170], radius=20, fill=rgb, outline=(15, 32, 67), width=3)
    draw.text((20, 180), caption, fill=(15, 32, 67))
    img.save(path)
    return path


def draw_cover_mosaic(path: Path):
    scale = 28
    img = Image.new("RGB", (3 * scale, 3 * scale))
    for i in range(3):
        for j in range(3):
            for dy in range(scale):
                for dx in range(scale):
                    img.putpixel((j * scale + dx, i * scale + dy), (R[i][j], G[i][j], B[i][j]))
    img = img.resize((360, 360), Image.NEAREST)
    img.save(path)
    return path


img_rgb = draw_grid_rgb(ASSETS / "grid_rgb.png", R, G, B)
img_r = draw_grid_gray(ASSETS / "grid_r.png", R, tint="R")
img_g = draw_grid_gray(ASSETS / "grid_g.png", G, tint="G")
img_b = draw_grid_gray(ASSETS / "grid_b.png", B, tint="B")
img_rt = draw_grid_gray(ASSETS / "grid_r_T.png", transpose(R), tint="R")
img_r05 = draw_grid_gray(ASSETS / "grid_r_05.png", scale_mat(R, 0.5), tint="R")
img_cover = draw_cover_mosaic(ASSETS / "cover_mosaic.png")
sw_r = draw_swatch(ASSETS / "swatch_r.png", (196, 40, 40), "R = 196")
sw_g = draw_swatch(ASSETS / "swatch_g.png", (40, 188, 40), "G = 188")
sw_b = draw_swatch(ASSETS / "swatch_b.png", (40, 80, 177), "B = 177")
sw_pix = draw_swatch(ASSETS / "swatch_pixel.png", (196, 188, 177), "pixel RGB")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = COLOR_BG


def header(slide, title: str):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.35)).text_frame
    p = t.paragraphs[0]
    p.text = "ALGEBRA LINEAR · CIENCIA DA COMPUTACAO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    t2 = slide.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(12), Inches(0.7)).text_frame
    p2 = t2.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NAVY


def card(slide, x, y, w, h, fill=COLOR_CARD, line=COLOR_LINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.25)
    return sh


def label(slide, x, y, w, text, size=14, bold=False, color=COLOR_DARK, align=PP_ALIGN.LEFT):
    tf = slide.shapes.add_textbox(x, y, w, Inches(0.4)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align


# SLIDE 1
s1 = prs.slides.add_slide(blank)
bg(s1)
card(s1, Inches(0.9), Inches(1.3), Inches(7.6), Inches(4.9), fill=COLOR_CARD, line=COLOR_BLUE)
tf = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(7.0), Inches(4.0)).text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Matrizes e sistemas lineares\nem imagens digitais"
p.font.size = Pt(34)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(18)
p2 = tf.add_paragraph()
p2.text = "Algebra Linear · Ciencia da Computacao"
p2.font.size = Pt(18)
p2.font.bold = True
p2.font.color.rgb = COLOR_BLUE
p2.space_after = Pt(12)
p3 = tf.add_paragraph()
p3.text = "Lucas Muller Scuzziato  ·  18/09/2026"
p3.font.size = Pt(15)
p3.font.color.rgb = COLOR_MUTED
p4 = tf.add_paragraph()
p4.text = "\nIdeia-chave: 3 matrizes alinhadas → 1 imagem."
p4.font.size = Pt(14)
p4.font.color.rgb = COLOR_MUTED
s1.shapes.add_picture(str(img_cover), Inches(9.0), Inches(1.8), Inches(3.5), Inches(3.5))
label(s1, Inches(9.0), Inches(5.45), Inches(3.5), "pixels = numeros", size=13, bold=True, color=COLOR_BLUE, align=PP_ALIGN.CENTER)

# SLIDE 2 — o que voce pediu
s2 = prs.slides.add_slide(blank)
bg(s2)
header(s2, "Imagem colorida = 3 matrizes alinhadas")

card(s2, Inches(0.55), Inches(1.4), Inches(12.2), Inches(1.35), fill=COLOR_SOFT, line=COLOR_BLUE)
tf = s2.shapes.add_textbox(Inches(0.75), Inches(1.5), Inches(11.8), Inches(1.15)).text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Ex.: foto 20×20 → matriz R 20×20 + matriz G 20×20 + matriz B 20×20"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p2 = tf.add_paragraph()
p2.text = "Mesma posicao (linha, coluna) nas tres = o mesmo pixel. Na tela, R+G+B daquela casa viram 1 cor."
p2.font.size = Pt(14)
p2.font.color.rgb = COLOR_DARK

trio = [
    (Inches(0.4), "R (vermelho)", img_r),
    (Inches(3.45), "G (verde)", img_g),
    (Inches(6.5), "B (azul)", img_b),
]
for x, title, img in trio:
    label(s2, x, Inches(2.9), Inches(2.9), title, size=14, bold=True, color=COLOR_BLUE, align=PP_ALIGN.CENTER)
    s2.shapes.add_picture(str(img), x + Inches(0.15), Inches(3.3), Inches(2.6), Inches(2.6))

label(s2, Inches(9.3), Inches(2.9), Inches(3.4), "→ cor na tela", size=14, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
s2.shapes.add_picture(str(img_rgb), Inches(9.45), Inches(3.3), Inches(2.9), Inches(2.9))
label(
    s2,
    Inches(0.55),
    Inches(6.85),
    Inches(12),
    "Abaixo: recorte real 3×3 (mesmo esquema). Nao e bagunca: sao 3 tabelas; a mistura e por pixel.",
    size=13,
    color=COLOR_MUTED,
)

# SLIDE 3
s3 = prs.slides.add_slide(blank)
bg(s3)
header(s3, "Operacoes de matriz → efeito visual (dia a dia)")

cols = [
    (Inches(0.55), "Original R", img_r, "matriz R"),
    (Inches(4.55), "Transposta R^T", img_rt, "linhas ↔ colunas"),
    (Inches(8.55), "Escalar 0,5·R", img_r05, "contraste ↓"),
]
for x, title, img, sub in cols:
    card(s3, x, Inches(1.5), Inches(3.85), Inches(5.5))
    label(s3, x + Inches(0.15), Inches(1.65), Inches(3.5), title, size=16, bold=True, color=COLOR_BLUE, align=PP_ALIGN.CENTER)
    label(s3, x + Inches(0.15), Inches(2.0), Inches(3.5), sub, size=12, color=COLOR_MUTED, align=PP_ALIGN.CENTER)
    s3.shapes.add_picture(str(img), x + Inches(0.35), Inches(2.45), Inches(3.15), Inches(3.15))

label(
    s3,
    Inches(0.7),
    Inches(7.05),
    Inches(12),
    "No codigo: img[y][x]. Transposta espelha; escalar muda o contraste.",
    size=13,
    color=COLOR_MUTED,
)

# SLIDE 4
s4 = prs.slides.add_slide(blank)
bg(s4)
header(s4, "Sistema linear (exemplo didatico)")

card(s4, Inches(0.6), Inches(1.5), Inches(5.9), Inches(5.5))
tf = s4.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(5.4), Inches(5.1)).text_frame
tf.word_wrap = True
lines = [
    ("Exemplo didatico (nao e o fluxo do dia a dia)", True, COLOR_NAVY, 15),
    ("R + G + B = 561", False, COLOR_DARK, 16),
    ("2R + G      = 580", False, COLOR_DARK, 16),
    ("R + 2B      = 550", False, COLOR_DARK, 16),
    ("", False, COLOR_DARK, 8),
    ("Forma A · X = B", True, COLOR_NAVY, 17),
    ("|1 1 1| |R|   |561|", False, COLOR_DARK, 14),
    ("|2 1 0| |G| = |580|", False, COLOR_DARK, 14),
    ("|1 0 2| |B|   |550|", False, COLOR_DARK, 14),
    ("", False, COLOR_DARK, 8),
    ("Gauss → B=177, G=188, R=196", True, COLOR_BLUE, 15),
]
first = True
for text, bold, color, size in lines:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.text = text
    p.font.bold = bold
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.name = "Consolas" if "|" in text or text.startswith("R ") or text.startswith("2R") else "Calibri"
    p.space_after = Pt(4)

card(s4, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.5), fill=COLOR_SOFT, line=COLOR_BLUE)
label(s4, Inches(7.0), Inches(1.7), Inches(5.5), "Resposta = 1 pixel (R,G,B)", size=16, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
s4.shapes.add_picture(str(sw_r), Inches(7.0), Inches(2.3), Inches(1.7), Inches(1.7))
s4.shapes.add_picture(str(sw_g), Inches(8.9), Inches(2.3), Inches(1.7), Inches(1.7))
s4.shapes.add_picture(str(sw_b), Inches(10.8), Inches(2.3), Inches(1.7), Inches(1.7))
s4.shapes.add_picture(str(sw_pix), Inches(8.5), Inches(4.3), Inches(2.4), Inches(2.4))
label(s4, Inches(7.0), Inches(6.55), Inches(5.5), "Conferencia: 2·196 + 188 = 580", size=13, bold=True, color=COLOR_BLUE, align=PP_ALIGN.CENTER)

# SLIDE 5
s5 = prs.slides.add_slide(blank)
bg(s5)
header(s5, "Conclusao")

items = [
    ("1", "3 matrizes alinhadas", "R, G e B → cor por pixel."),
    ("2", "Operacoes = efeitos", "Transposta e escalar no dia a dia."),
    ("3", "Sistema = ferramenta", "Exemplo didatico com Gauss."),
]
for i, (n, title, desc) in enumerate(items):
    x = Inches(0.7) + i * Inches(4.15)
    card(s5, x, Inches(1.7), Inches(3.9), Inches(3.6), line=COLOR_BLUE if i == 0 else COLOR_LINE)
    previews = [img_rgb, img_rt, sw_pix]
    s5.shapes.add_picture(str(previews[i]), x + Inches(0.95), Inches(1.95), Inches(2.0), Inches(2.0))
    label(s5, x + Inches(0.2), Inches(4.1), Inches(3.5), f"{n}. {title}", size=16, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
    label(s5, x + Inches(0.2), Inches(4.55), Inches(3.5), desc, size=13, color=COLOR_MUTED, align=PP_ALIGN.CENTER)

label(
    s5,
    Inches(0.7),
    Inches(5.7),
    Inches(12),
    "STEINBRUCH; WINTERLE — Algebra linear (2014)  ·  GONZALEZ; WOODS — Processamento digital de imagens",
    size=12,
    color=COLOR_MUTED,
)
label(s5, Inches(0.7), Inches(6.3), Inches(12), "Obrigado.", size=22, bold=True, color=COLOR_NAVY)

out = OUT_DIR / "Matrizes_e_Sistemas_Lineares_Imagens_VISUAL.pptx"
try:
    prs.save(str(out))
except PermissionError:
    out = OUT_DIR / "Matrizes_e_Sistemas_Lineares_Imagens_VISUAL_v2.pptx"
    prs.save(str(out))
print(f"OK -> {out}")
